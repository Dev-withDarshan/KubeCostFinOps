"""
assemble_and_generate.py
Main entry point to assemble all report sections and write out the final .docx report.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_report_sections import (
    build_front_matter,
    build_chapter_1,
    build_chapter_2,
    build_chapter_3,
    build_chapter_4,
    build_chapter_5
)

from build_report_sections_part2 import (
    build_chapter_6,
    build_chapter_7,
    build_chapter_8,
    build_chapter_9,
    build_chapter_10
)

from build_report_sections_part3 import (
    build_chapter_11,
    build_chapter_12,
    build_chapter_13,
    build_chapter_14,
    build_references,
    build_appendices
)

def setup_page_layout_and_footer(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    
    # Running Header
    hdr = sec.header
    p_hdr = hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run("BCSE355L — Cloud Architecture Design  |  Final Semester Project Report")
    r_hdr.font.name = "Times New Roman"
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.color.rgb = RGBColor(140, 150, 160)
    
    # Running Footer with Page Number XML
    ftr = sec.footer
    p_ftr = ftr.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr1 = p_ftr.add_run("Kubernetes Cost Optimization Framework  |  Page ")
    r_ftr1.font.name = "Times New Roman"
    r_ftr1.font.size = Pt(9.5)
    r_ftr1.font.color.rgb = RGBColor(100, 116, 139)
    
    r_ftr2 = p_ftr.add_run()
    r_ftr2.font.name = "Times New Roman"
    r_ftr2.font.size = Pt(9.5)
    r_ftr2.font.color.rgb = RGBColor(100, 116, 139)
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = " PAGE "
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'end')
    r_ftr2._r.append(fld1)
    r_ftr2._r.append(instr)
    r_ftr2._r.append(fld2)

def generate_report():
    print("Creating Document instance...")
    doc = Document()
    setup_page_layout_and_footer(doc)
    
    print("Building Section 0: Front Matter (Cover, Abstract, TOC)...")
    build_front_matter(doc)
    
    print("Building Chapter 1: Introduction...")
    build_chapter_1(doc)
    
    print("Building Chapter 2: Literature Survey / Related Work (Table 0 & Table 1)...")
    build_chapter_2(doc)
    
    print("Building Chapter 3: Research Gap Analysis...")
    build_chapter_3(doc)
    
    print("Building Chapter 4: Project Objectives...")
    build_chapter_4(doc)
    
    print("Building Chapter 5: Novelty Summary (Table 2)...")
    build_chapter_5(doc)
    
    print("Building Chapter 6: Proposed System and Architecture (Figures 1, 2, 3)...")
    build_chapter_6(doc)
    
    print("Building Chapter 7: Dataset Description (Table 3)...")
    build_chapter_7(doc)
    
    print("Building Chapter 8: AWS Services Planning (Table 4)...")
    build_chapter_8(doc)
    
    print("Building Chapter 9: Methodology and Implementation Plan (Table 5)...")
    build_chapter_9(doc)
    
    print("Building Chapter 10: Implementation / Development (Table 6)...")
    build_chapter_10(doc)
    
    print("Building Chapter 11: Testing, Results and Observations (Table 7, Figures 4, 5, 6)...")
    build_chapter_11(doc)
    
    print("Building Chapter 12: Discussion, Limitations and Risk Analysis (Table 8)...")
    build_chapter_12(doc)
    
    print("Building Chapter 13: Team Contribution and GitHub Management (Table 9)...")
    build_chapter_13(doc)
    
    print("Building Chapter 14: Conclusion and Future Enhancements...")
    build_chapter_14(doc)
    
    print("Building References Section...")
    build_references(doc)
    
    print("Building Appendices (Timeline Table 10, Review Checklist, Reproducibility)...")
    build_appendices(doc)
    
    # Save outputs
    out_paths = [
        "1Submission/BCSE355L_Kubernetes_FinOps_Final_Report.docx",
        "BCSE355L_Kubernetes_FinOps_Final_Report.docx",
        "1Submission/FINAL_VERIFIED/BCSE355L_Kubernetes_FinOps_Final_Report.docx",
        "1Submission/BCSE355L_Final_Project_Report_K_Darshan.docx"
    ]
    for p in out_paths:
        os.makedirs(os.path.dirname(os.path.abspath(p)), exist_ok=True)
        doc.save(p)
        print(f"Successfully generated: {p} ({os.path.getsize(p)} bytes)")
        
    print("\nReport generation completed successfully!")

if __name__ == "__main__":
    generate_report()
