"""
build_full_report.py
Generates the comprehensive, publication-grade Final Semester Project Report (.docx)
for BCSE355L - Cloud Architecture Design:
'Kubernetes Cost Optimization Framework for SaaS-Based Educational Platforms using FinOps Intelligence'

Follows the instructor's report template (AWS_Cloud_Project_Report_Template.docx) exactly,
carrying forward and updating the Phase-I / Review 1 foundation with empirical evidence,
regression test results, full 15-paper comparative matrix, and rigorous academic analysis.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Import structured data
from report_data import (
    PROJECT_METADATA,
    PAPERS_DATA,
    NOVELTY_DATA,
    DATASET_ATTRIBUTES,
    AWS_SERVICES_DATA,
    TOOLS_DATA,
    MODULES_DATA,
    TEST_CASES_DATA,
    RISKS_DATA,
    TEAM_CONTRIBUTION_DATA,
    TIMELINE_DATA,
    CHECKLIST_ITEMS
)

print("Starting full final report generation...")

# ---------------------------------------------------------------------------
# Helper XML Formatting Functions
# ---------------------------------------------------------------------------
def set_cell_shading(cell, color_hex):
    """Apply background color fill to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set inner cell padding in twips/dxa (20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Apply clean, professional subtle borders to an entire table."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def style_header_cell(cell, text, width=None):
    """Format table header cells with dark navy fill, white bold text."""
    set_cell_shading(cell, "1F4E79")
    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    if width:
        cell.width = width

def style_body_cell(cell, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, width=None, bg_color=None):
    """Format regular table body cells."""
    if bg_color:
        set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.0)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(30, 41, 59)
    if width:
        cell.width = width

def setup_table_behavior(table):
    """Configure table to repeat header row and avoid row breaking across pages."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in table.rows:
        trPr = r._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    if len(table.rows) > 0:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

# ---------------------------------------------------------------------------
# Paragraph & Document Helpers
# ---------------------------------------------------------------------------
def add_p(doc, text="", style="Normal", space_before=0, space_after=6, line_spacing=1.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121) # Professional Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(46, 64, 83)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout_box(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_shading(c, "F8FAFC") # Slate 50
    set_cell_margins(c, top=100, bottom=100, left=160, right=160)
    # Left border only (accent border)
    tcPr = c._tc.get_or_add_tcPr()
    tblBorders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 3pt
    left.set(qn('w:color'), '1F4E79')
    tblBorders.append(left)
    for b in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tcPr.append(tblBorders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(title + "\n")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 78, 121)
    
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    # Add empty spacing paragraph after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_image_with_caption(doc, image_path, caption_text, width=Inches(5.8)):
    if not os.path.exists(image_path):
        p_warn = add_p(doc, f"[Image placeholder: {os.path.basename(image_path)} not found]", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = True
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

print("Document and paragraph helpers ready.")
